# -*- coding: utf-8 -*-
"""DOCX 排版工具：中文报告样式"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pandas as pd

BODY_CN = "宋体"
HEAD_CN = "黑体"
LATIN = "Times New Roman"
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x2E, 0x5E, 0x8E)



PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
             "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]

TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
              "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
              "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"]

TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
              "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
              "ins", "del", "trPrChange"]

TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
               "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription", "tblPrChange"]


def _ordered_insert(parent, child, order):
    """按 OOXML schema 顺序插入子元素，覆盖同名的旧元素"""
    from lxml import etree
    tag = etree.QName(child).localname
    for old in parent.findall(child.tag):
        parent.remove(old)
    idx = order.index(tag) if tag in order else len(order)
    for existing in parent:
        etag = etree.QName(existing).localname
        if etag in order and order.index(etag) > idx:
            existing.addprevious(child)
            return
    parent.append(child)


def _border(parent, edge="bottom", color="2E5E8E", sz="8", space="1", order=None):
    from lxml import etree
    tag = "{%s}pBdr" % etree.QName(parent).namespace
    pbdr = parent.find(tag)
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        _ordered_insert(parent, pbdr, order or PPR_ORDER)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), space)
    el.set(qn("w:color"), color)
    pbdr.append(el)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)
    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    pf = st.paragraph_format
    pf.line_spacing = 1.45
    pf.space_after = Pt(2)
    return doc


def _set_run(run, size=10.5, bold=False, color=None, cn=BODY_CN, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = LATIN
    run.font.color.rgb = color if color is not None else DARK
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:eastAsia"), cn)
    rf.set(qn("w:ascii"), LATIN)
    rf.set(qn("w:hAnsi"), LATIN)


def title(doc, text, sub=None, meta=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(6)
    _set_run(p.add_run(text), 20, True, DARK, HEAD_CN)
    if sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(4)
        _set_run(p2.add_run(sub), 12, False, GREY, BODY_CN)
    if meta:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_before = Pt(10)
        _set_run(p3.add_run(meta), 9.5, False, GREY, BODY_CN)
    # 分隔线
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_before = Pt(10)
    p4.paragraph_format.space_after = Pt(14)
    pPr = p4._p.get_or_add_pPr()
    _border(pPr, sz="8", color="2E5E8E")


def h1(doc, text, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    _set_run(p.add_run(text), 14.5, True, ACCENT, HEAD_CN)
    pPr = p._p.get_or_add_pPr()
    _border(pPr, sz="6", color="C9D6E4", space="3")
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _set_run(p.add_run(text), 11.5, True, DARK, HEAD_CN)
    return p


def para(doc, text, indent=True, size=10.5, space_after=3, align=None, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.45
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify" or align is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_run(p.add_run(text), size, False, color or DARK)
    return p


def rich(doc, parts, indent=True, size=10.5, space_after=3):
    """parts: [(text, bold), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.45
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t, b in parts:
        _set_run(p.add_run(t), size, b)
    return p


def bullets(doc, items, size=10.5, indent_pt=21, space_after=2):
    for it in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(indent_pt + 12)
        pf.first_line_indent = Pt(-12)
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.4
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _set_run(p.add_run("· "), size, False, ACCENT)
        _set_run(p.add_run(it), size, False, DARK)


def numbered(doc, items, size=10.5, indent_pt=21):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(indent_pt + 16)
        pf.first_line_indent = Pt(-16)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.4
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _set_run(p.add_run(f"（{i}）"), size, False, ACCENT)
        _set_run(p.add_run(it), size, False, DARK)


def caption(doc, text, above=False, size=9, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4 if not above else 8)
    p.paragraph_format.space_after = Pt(3 if note else (8 if not above else 3))
    p.paragraph_format.keep_with_next = above
    _set_run(p.add_run(text), size, True, GREY, HEAD_CN)
    if note:
        pn = doc.add_paragraph()
        pn.paragraph_format.space_before = Pt(1)
        pn.paragraph_format.space_after = Pt(8)
        pn.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_run(pn.add_run(note), size - 1, False, GREY)
    return p


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    _ordered_insert(tcPr, shd, TCPR_ORDER)


def _cell_borders(cell, color="B8C4D2", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    _ordered_insert(tcPr, borders, TCPR_ORDER)


def table(doc, df, widths=None, font_size=8.5, header_size=8.5, first_col_bold=False,
          align_num=True, row_height=None, max_rows=None, note=None):
    df = df.copy()
    if max_rows and len(df) > max_rows:
        df = df.head(max_rows)
    ncol = len(df.columns)
    nrow = len(df) + 1
    t = doc.add_table(rows=nrow, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    # 表头跨页重复 + 行不拆分
    for i, r in enumerate(t.rows):
        pr = r._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        _ordered_insert(pr, cs, TRPR_ORDER)
        if i == 0:
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
            _ordered_insert(pr, th, TRPR_ORDER)
    # header
    for j, col in enumerate(df.columns):
        c = t.cell(0, j)
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        _set_run(p.add_run(str(col)), header_size, True, RGBColor(0x1F, 0x3A, 0x55), HEAD_CN)
        _shade(c, "E4EBF3")
        _cell_borders(c)
        c.vertical_alignment = 1
    # body
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(df.columns):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            val = row[col]
            if isinstance(val, float):
                val = f"{val:,.2f}".rstrip("0").rstrip(".") if abs(val) < 1000 else f"{val:,.0f}"
            s = "" if pd.isna(val) else str(val)
            is_num = align_num and j > 0 and s.replace(",", "").replace(".", "").replace("%", "").replace("-", "").isdigit()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_num else (
                WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(1.2)
            p.paragraph_format.space_after = Pt(1.2)
            p.paragraph_format.line_spacing = 1.0
            bold = first_col_bold and j == 0
            _set_run(p.add_run(s), font_size, bold, DARK)
            _cell_borders(c)
            c.vertical_alignment = 1
        if row_height:
            t.rows[i].height = row_height
    if widths:
        tot = sum(widths)
        if tot > 15.8:
            widths = [w * 15.8 / tot for w in widths]
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Cm(w)
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        _set_run(p.add_run(note), 8, False, GREY)
    return t


def figure(doc, path, width_cm=15.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))


def add_page_number_footer(doc, left_text=""):
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if left_text:
        _set_run(p.add_run(left_text + "     "), 8, False, GREY)
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    _set_run(run, 8, False, GREY)


def add_header_text(doc, text):
    sec = doc.sections[0]
    p = sec.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run(p.add_run(text), 8, False, GREY)
    pPr = p._p.get_or_add_pPr()
    _border(pPr, sz="4", color="D9D9D9")


def keybox(doc, lines, title_text=None):
    """要点框"""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    c.text = ""
    _shade(c, "F2F5F9")
    _cell_borders(c, color="C9D6E4", sz="6")
    tcPr = c._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", "120"), ("bottom", "120"), ("left", "180"), ("right", "180")):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), v); el.set(qn("w:type"), "dxa")
        mar.append(el)
    _ordered_insert(tcPr, mar, TCPR_ORDER)
    first = True
    if title_text:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(3)
        _set_run(p.add_run(title_text), 9.5, True, ACCENT, HEAD_CN)
        first = False
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.left_indent = Pt(10)
        p.paragraph_format.first_line_indent = Pt(-10)
        _set_run(p.add_run("· "), 9, False, ACCENT)
        _set_run(p.add_run(ln), 9, False, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
